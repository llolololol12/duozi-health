import openpyxl, zipfile, json, os, re

XLSX = r'C:\Users\20461\Desktop\Product Information产品信息汇总表2026.0708.xlsx'
OUT_DIR = r'C:\Users\20461\Documents\ChatGPT\产品页\public\images\products'
os.makedirs(OUT_DIR, exist_ok=True)

z = zipfile.ZipFile(XLSX)
cell_xml = z.read('xl/cellimages.xml').decode('utf-8', errors='replace')
id2rid = {m.group(1): m.group(2) for m in re.finditer(r'name="(ID_[A-F0-9]+)"[^>]*>.*?<a:blip r:embed="(rId\d+)"', cell_xml, re.S)}
rels_xml = z.read('xl/_rels/cellimages.xml.rels').decode('utf-8', errors='replace')
rid2media = {m.group(1): m.group(2) for m in re.finditer(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', rels_xml)}
id2media = {iid: 'xl/media/' + rid2media[rid].replace('\\','/').split('/')[-1] for iid,rid in id2rid.items() if rid in rid2media}

def txt(v):
    return '' if v is None else str(v).strip()

def safe_code(v):
    if v is None:
        return None
    s = re.sub(r'[^A-Za-z0-9_-]+', '', str(v).strip()).upper()
    return s or None

wb = openpyxl.load_workbook(XLSX, data_only=True, read_only=True)
ws = wb[wb.sheetnames[0]]
rows = list(ws.iter_rows(values_only=True))
hdr = next(i for i,r in enumerate(rows) if r and r[0] == 'NO')

products = []
missing_img = 0
total_bytes = 0
used_codes = set()
for idx, r in enumerate(rows[hdr+1:]):
    if not r or r[0] is None:
        continue
    code = safe_code(r[0]) or f'P{idx+1}'
    base = code; k = 2
    while code in used_codes:
        code = f'{base}-{k}'; k += 1
    used_codes.add(code)
    formula = txt(r[1])
    m = re.search(r'ID_([A-F0-9]+)', formula)
    media = id2media.get('ID_' + m.group(1)) if m else None
    image_path = None
    if media and media in z.namelist():
        data = z.read(media)
        total_bytes += len(data)
        ext = os.path.splitext(media)[1]
        fname = f'product-{code}{ext}'
        with open(os.path.join(OUT_DIR, fname), 'wb') as f:
            f.write(data)
        image_path = '/images/products/' + fname
    else:
        missing_img += 1
    name = txt(r[2])
    en_name = name.split('\n')[0].strip() if name else ''
    products.append({
        'code': code,
        'name': en_name,
        'name_full': name,
        'type': txt(r[4]),
        'spec': txt(r[3]),
        'net': txt(r[5]),
        'gross': txt(r[6]),
        'ctn': txt(r[7]),
        'function': txt(r[8]),
        'competitor': txt(r[9]),
        'image': image_path,
    })

json.dump(products, open(r'C:\Users\20461\Documents\ChatGPT\产品页\.excel-mapped.json','w',encoding='utf-8'), ensure_ascii=False, indent=1)
n = len(products)
nimg = sum(1 for p in products if p['image'])
print('products:', n)
print('with image:', nimg)
print('missing image:', missing_img)
print('total image bytes: %.1f MB' % (total_bytes/1024/1024))
print('avg image KB: %.1f' % (total_bytes/1024/max(1,nimg)))
